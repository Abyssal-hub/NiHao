#!/usr/bin/env python3
"""Extract text from all HSK1 PPTX files."""
import os
from pptx import Presentation
from glob import glob

downloads_dir = "/root/.openclaw/workspace/downloads"
output_dir = "/root/.openclaw/workspace"

pptx_files = sorted(glob(os.path.join(downloads_dir, "*.pptx")))

with open(os.path.join(output_dir, "hsk1_extracted.txt"), "w", encoding="utf-8") as out:
    for pptx_path in pptx_files:
        filename = os.path.basename(pptx_path)
        out.write(f"\n{'='*60}\n")
        out.write(f"FILE: {filename}\n")
        out.write(f"{'='*60}\n\n")
        
        try:
            prs = Presentation(pptx_path)
            for i, slide in enumerate(prs.slides, 1):
                out.write(f"--- Slide {i} ---\n")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        out.write(shape.text.strip() + "\n")
                out.write("\n")
        except Exception as e:
            out.write(f"ERROR: {e}\n\n")

print(f"Extracted text from {len(pptx_files)} files to hsk1_extracted.txt")
